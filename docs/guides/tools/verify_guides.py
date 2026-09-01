#!/usr/bin/env python3
"""Verify cross-format identity and publication hygiene for BSEW guides."""

from __future__ import annotations

import argparse
import re
import subprocess
import sys
from html import unescape
from html.parser import HTMLParser
from pathlib import Path

from docx import Document


LOCAL_PATHS = re.compile(r"(?:/Users/|/Volumes/|file://|[A-Za-z]:\\Users\\)")
OLD_VISIBLE_NAMES = (
    "Model Builder Research Workbench",
    "Valencia Energy Simulation Workbench",
    "Valencia Workbench",
    "Academic User Guide",
    "ACADEMIC RESEARCH DOCUMENTATION",
)
KNOWN_BROKEN_LINKS = (
    "https://energyplus.net/documentation",
)


class GuideHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.headings: list[tuple[int, str]] = []
        self.ids: list[str] = []
        self.captions: list[str] = []
        self.figure_captions: list[str] = []
        self.image_alts: list[str | None] = []
        self.th_scopes: list[str | None] = []
        self.lang: str | None = None
        self._capture: tuple[str, list[str]] | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        values = dict(attrs)
        if tag == "html":
            self.lang = values.get("lang")
        if values.get("id"):
            self.ids.append(values["id"] or "")
        if tag in {"h1", "h2", "h3", "h4", "caption", "figcaption"}:
            self._capture = (tag, [])
        if tag == "img":
            self.image_alts.append(values.get("alt"))
        if tag == "th":
            self.th_scopes.append(values.get("scope"))

    def handle_data(self, data: str) -> None:
        if self._capture:
            self._capture[1].append(data)

    def handle_endtag(self, tag: str) -> None:
        if self._capture and self._capture[0] == tag:
            text = " ".join("".join(self._capture[1]).split())
            if tag.startswith("h"):
                self.headings.append((int(tag[1]), text))
            elif tag == "caption":
                self.captions.append(text)
            else:
                self.figure_captions.append(text)
            self._capture = None


def normalise(text: str) -> str:
    collapsed = " ".join(text.replace("\u00a0", " ").split())
    return re.sub(r"\s+([.,;:])", r"\1", collapsed)


def plain_markdown(text: str) -> str:
    return normalise(re.sub(r"[`*]", "", text))


def source_contract(path: Path) -> tuple[str, list[str], int, int, int, str]:
    text = path.read_text(encoding="utf-8")
    cover, body = text.split("<!-- pagebreak -->", 1)
    cover_headings = [match.group(2).strip() for match in re.finditer(r"^(#{1,4})\s+(.+)$", cover, re.MULTILINE)]
    title = cover_headings[0]
    body_headings = [match.group(2).strip() for match in re.finditer(r"^(#{1,4})\s+(.+)$", body, re.MULTILINE)]
    table_count = len(re.findall(r"^\|.+\|$\n^\|\s*:?-", text, re.MULTILINE))
    figure_count = len(re.findall(r'^!\[[^]]+\]\(\S+(?:\s+"[^"]+")?\)$', text, re.MULTILINE))
    reference_count = len(re.findall(r"^\[\d+\]\s", text, re.MULTILINE))
    citation_match = re.search(r"^### Suggested citation\s*$\n\n(.+)$", text, re.MULTILINE)
    if not citation_match:
        raise ValueError(f"Suggested citation not found in {path}")
    return title, body_headings, table_count, figure_count, reference_count, plain_markdown(citation_match.group(1))


def source_hygiene(path: Path) -> list[str]:
    """Catch publication errors that format-parity checks cannot detect."""
    errors: list[str] = []
    text = path.read_text(encoding="utf-8")
    reference_numbers = {
        int(value) for value in re.findall(r"^\[(\d+)\]\s", text, re.MULTILINE)
    }
    reference_heading = re.search(r"^#{1,2}\s+(?:\d+\.\s+)?References(?:\s+and document control)?\s*$", text, re.MULTILINE)
    body = text[: reference_heading.start()] if reference_heading else text
    cited_numbers = {int(value) for value in re.findall(r"\[(\d+)\]", body)}
    for start, end in re.findall(r"\[(\d+)\]\s*[–-]\s*\[(\d+)\]", body):
        cited_numbers.update(range(int(start), int(end) + 1))
    uncited = sorted(reference_numbers - cited_numbers)
    if uncited:
        errors.append(f"numbered references are not cited in the text: {uncited}")
    if "Appendix A" in text and not re.search(r"^#{1,4}\s+Appendix A\b", text, re.MULTILINE):
        errors.append("text refers to Appendix A, but no Appendix A heading exists")
    for link in KNOWN_BROKEN_LINKS:
        if link in text:
            errors.append(f"source contains a known broken link: {link}")
    if "`py list`" in text:
        errors.append("Windows guidance uses py list instead of the legacy-compatible py --list")
    if path.name == "user-guide.md":
        if "refuses a bare `.shp`" not in text:
            errors.append("Files guidance must state that a bare .shp upload is refused")
        if "more than one `.shp` dataset" not in text:
            errors.append("Files guidance must state the single-Shapefile-per-ZIP rule")
    return errors


def html_contract(path: Path) -> tuple[GuideHTMLParser, str]:
    text = path.read_text(encoding="utf-8")
    parser = GuideHTMLParser()
    parser.feed(text)
    return parser, text


def docx_contract(path: Path) -> tuple[list[str], str, int, int]:
    document = Document(path)
    headings = [normalise(paragraph.text) for paragraph in document.paragraphs if paragraph.style.name.startswith("Heading")]
    all_text = normalise("\n".join(paragraph.text for paragraph in document.paragraphs))
    return headings, all_text, len(document.tables), len(document.inline_shapes)


def pdf_text(path: Path) -> str:
    result = subprocess.run(
        ["pdftotext", "-layout", str(path), "-"],
        check=True,
        capture_output=True,
        text=True,
    )
    return normalise(result.stdout)


def verify_one(source: Path, html_path: Path, docx_path: Path, pdf_path: Path, final: bool) -> list[str]:
    missing = [path for path in (source, html_path, docx_path, pdf_path) if not path.is_file()]
    if missing:
        return ["expected publication artifact is missing: " + ", ".join(path.name for path in missing)]
    errors: list[str] = source_hygiene(source)
    title, expected_headings, expected_tables, expected_figures, expected_refs, citation = source_contract(source)
    html_parser, html_text = html_contract(html_path)
    html_visible_text = normalise(unescape(re.sub(r"<[^>]+>", " ", html_text)))
    docx_headings, docx_text, docx_tables, docx_figures = docx_contract(docx_path)
    rendered_pdf_text = pdf_text(pdf_path)

    html_heading_text = [text for _, text in html_parser.headings]
    if len([1 for level, _ in html_parser.headings if level == 1]) != 1:
        errors.append("HTML must contain exactly one H1")
    if html_heading_text[-len(expected_headings):] != expected_headings:
        errors.append("HTML body headings differ from canonical source")
    if docx_headings != expected_headings:
        errors.append("DOCX headings differ from canonical source")
    if len(html_parser.ids) != len(set(html_parser.ids)):
        errors.append("HTML contains duplicate IDs")
    if html_parser.lang != "en":
        errors.append("HTML language is not en")
    if len(html_parser.captions) != expected_tables:
        errors.append("HTML table caption count differs from canonical source")
    if docx_tables != expected_tables:
        errors.append("DOCX data-table count differs from canonical source")
    if len(html_parser.figure_captions) != expected_figures:
        errors.append("HTML figure count differs from canonical source")
    if docx_figures != expected_figures:
        errors.append("DOCX figure count differs from canonical source")
    if len(html_parser.image_alts) != expected_figures or any(not item for item in html_parser.image_alts):
        errors.append("HTML figures must all carry non-empty alt text")
    if not html_parser.th_scopes or any(scope not in {"row", "col"} for scope in html_parser.th_scopes):
        errors.append("HTML table headers must all carry row/col scope")

    for label, text in (
        ("source", plain_markdown(source.read_text(encoding="utf-8"))),
        ("HTML", html_visible_text),
        ("DOCX", docx_text),
        ("PDF", rendered_pdf_text),
    ):
        if LOCAL_PATHS.search(text):
            errors.append(f"{label} contains an absolute local path")
        for old_name in OLD_VISIBLE_NAMES:
            if old_name in text:
                errors.append(f"{label} contains old visible product name: {old_name}")
        if citation not in normalise(text):
            errors.append(f"{label} does not contain the canonical suggested citation")
        refs = len(re.findall(r"\[\d+\]", text))
        if refs < expected_refs:
            errors.append(f"{label} exposes fewer than {expected_refs} numbered references")
        if final and re.search(
            r"\b(?:provisional|draft|pending|final evidence slot|final figure|capture pending|publication edition)\b",
            text,
            re.IGNORECASE,
        ):
            errors.append(f"{label} contains a provisional marker in final mode")
    if title not in html_text or title not in docx_text or title not in rendered_pdf_text:
        errors.append("Document title is not consistent across HTML, DOCX and PDF")
    return errors


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--final", action="store_true")
    parser.add_argument("--publication-date", default="2026-08-26")
    args = parser.parse_args()
    suffix = "" if args.final else f" (Provisional {args.publication_date})"
    specs = (
        (
            "installation-guide.md",
            "installation-guide.html",
            f"Building Stock Energy Workbench - Installation Guide{suffix}.docx",
            f"Building Stock Energy Workbench - Installation Guide{suffix}.pdf",
        ),
        (
            "user-guide.md",
            "user-guide.html",
            f"Building Stock Energy Workbench - User Guide{suffix}.docx",
            f"Building Stock Energy Workbench - User Guide{suffix}.pdf",
        ),
        (
            "valencia-simulation-report.md",
            "valencia-simulation-report.html",
            f"Building Stock Energy Workbench - Valencia Simulation Report{suffix}.docx",
            f"Building Stock Energy Workbench - Valencia Simulation Report{suffix}.pdf",
        ),
    )
    errors: list[str] = []
    for source, html_name, docx, pdf in specs:
        current = verify_one(
            args.source_dir / source,
            args.output_dir / html_name,
            args.output_dir / docx,
            args.output_dir / pdf,
            args.final,
        )
        errors.extend(f"{source}: {error}" for error in current)
    if errors:
        print("GUIDE_VERIFICATION_FAILED")
        print("\n".join(f"- {error}" for error in errors))
        return 1
    print("GUIDE_VERIFICATION_OK")
    return 0


if __name__ == "__main__":
    sys.exit(main())
